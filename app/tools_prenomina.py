from __future__ import annotations

import math
import os
import re
import unicodedata
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from agents import function_tool
except ImportError:  # Allows deterministic local tests before installing openai-agents.
    def function_tool(func=None, **_kwargs):  # type: ignore[no-redef]
        if func is None:
            return lambda wrapped: wrapped
        return func


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT_PATH = PROJECT_ROOT / "output" / "prenomina_resultado.xlsx"
DEFAULT_FONDO_AHORRO_FACTOR = 0.11
DEFAULT_UMA_DIARIA = 117.31
DEFAULT_FONDO_AHORRO_TOPE_MODE = "mensual"
DEFAULT_USE_EXCEL_FONDO_AHORRO = True
DEFAULT_DIAS_BASE_PERIODO = 30.4
DEFAULT_DIAS_MES = 30.4
FONDO_AHORRO_TOPE_MODES = {"none", "mensual", "quincenal", "proporcional"}

REQUIRED_CALC_COLUMNS = [
    "empleado_id",
    "nombre",
    "fecha",
    "horas_trabajadas",
    "sueldo_hora",
    "bono",
    "descuento",
]

COLUMN_ALIASES = {
    "numero": "numero",
    "no": "numero",
    "codigo": "empleado_id",
    "cod": "empleado_id",
    "empleado id": "empleado_id",
    "empleado_id": "empleado_id",
    "id empleado": "empleado_id",
    "nombre": "nombre",
    "nombre completo": "nombre",
    "empleado": "nombre",
    "fecha": "fecha",
    "fecha asistencia": "fecha",
    "dia": "fecha",
    "horas": "horas_trabajadas",
    "horas trabajadas": "horas_trabajadas",
    "horas_trabajadas": "horas_trabajadas",
    "sueldo hora": "sueldo_hora",
    "sueldo_hora": "sueldo_hora",
    "salario hora": "sueldo_hora",
    "salario por hora": "sueldo_hora",
    "bono": "bono",
    "bonos": "bono",
    "descuento": "descuento",
    "descuentos": "descuento",
    "observaciones": "observaciones",
}

ADMIN_REQUIRED_COLUMNS = [
    "nombre_completo",
    "sueldo_nominal",
    "puntualidad",
    "asistencia",
    "vales_despensa",
    "fondo_ahorro",
    "percepcion_sueldos",
    "honorarios_asimilados",
    "gasolina_sueldo",
    "socio",
    "efectivo",
    "facturado",
    "deuda_carro",
    "sueldo_bruto_mensual",
    "sueldo_bruto_quincenal",
    "descuento",
    "sueldo_bruto_mensual_final",
    "sueldo_bruto_quincenal_final",
    "observaciones",
]

ADMIN_BASE_ALIASES = {
    "no": "numero",
    "cod": "codigo",
    "codigo": "codigo",
    "empresa": "empresa",
    "nombre completo": "nombre_completo",
    "area": "area",
    "departamento": "departamento",
    "puesto": "puesto",
    "lugar de trabajo": "lugar_trabajo",
    "fecha de ingreso": "fecha_ingreso",
    "anos de labores": "anos_labores",
    "fecha de baja": "fecha_baja",
    "cuenta con fondo de ahorro": "cuenta_fondo_ahorro",
    "salario diario": "salario_diario",
    "factor integracion imss": "factor_integracion_imss",
    "salario diario integrado": "salario_diario_integrado",
    "sueldo nominal": "sueldo_nominal",
    "puntualidad": "puntualidad",
    "asistencia": "asistencia",
    "vales de despensa": "vales_despensa",
    "fondo de ahorro": "fondo_ahorro",
    "percepcion sueldos": "percepcion_sueldos",
    "honorarios asimilados": "honorarios_asimilados",
    "gasolina sueldo": "gasolina_sueldo",
    "socio": "socio",
    "efectivo": "efectivo",
    "facturado": "facturado",
    "deuda de carro": "deuda_carro",
    "deude de carro": "deuda_carro",
    "sueldo bruto mensual": "sueldo_bruto_mensual",
    "sueldo bruto quincenal": "sueldo_bruto_quincenal",
    "sueldo bruto mensual final": "sueldo_bruto_mensual_final",
    "sueldo bruto quincenal final": "sueldo_bruto_quincenal_final",
    "porcentaje": "porcentaje",
    "descuento": "descuento",
    "dias trabajados": "dias_trabajados",
    "dias trabajados quincena": "dias_trabajados",
    "dias del mes": "dias_mes",
    "observaciones": "observaciones",
    "descuento empleados s quincenal": "descuento_empleados_quincenal",
    "descuento empleados s mensual": "descuento_empleados_mensual",
}


def _normalize_text(value: Any) -> str:
    text = "" if pd.isna(value) else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.lower().strip()
    text = re.sub(r"[\(\)\[\]\.:;,_/-]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _env_float(name: str, default: float) -> float:
    value = os.getenv(name)
    if value is None or value.strip() == "":
        return default
    return float(value)


def _env_bool(name: str, default: bool) -> bool:
    value = os.getenv(name)
    if value is None or value.strip() == "":
        return default
    normalized = _normalize_text(value)
    if normalized in {"1", "true", "yes", "si"}:
        return True
    if normalized in {"0", "false", "no"}:
        return False
    return default


def _normalize_yes_no(value: Any) -> bool:
    normalized = _normalize_text(value)
    if normalized in {"si", "yes", "true", "1"}:
        return True
    if normalized in {"", "no", "false", "0"}:
        return False
    return False


def _calculate_fondo_ahorro_cap(
    *,
    uma_diaria: float,
    tope_mode: str,
    dias_mes: float | None = None,
    dias_trabajados: float | None = None,
) -> dict[str, Any]:
    normalized_mode = _normalize_text(tope_mode) or DEFAULT_FONDO_AHORRO_TOPE_MODE
    if normalized_mode not in FONDO_AHORRO_TOPE_MODES:
        normalized_mode = DEFAULT_FONDO_AHORRO_TOPE_MODE

    monthly_cap = 1.3 * uma_diaria * 365 / 12
    warning = ""
    if normalized_mode == "none":
        cap = None
    elif normalized_mode == "mensual":
        cap = monthly_cap
    elif normalized_mode == "quincenal":
        cap = 1.3 * uma_diaria * 365 / 24
    else:
        if dias_mes and dias_mes > 0 and dias_trabajados and dias_trabajados > 0:
            cap = (monthly_cap / dias_mes) * dias_trabajados
        else:
            cap = monthly_cap
            normalized_mode = "mensual"
            warning = (
                "No hay dias_mes o dias_trabajados confiables para tope proporcional; "
                "se uso tope mensual como fallback."
            )

    return {
        "tope": round(cap, 2) if cap is not None else None,
        "tope_mode_usado": normalized_mode,
        "advertencia": warning,
    }


def _calculate_fondo_ahorro(
    *,
    cuenta_fondo_ahorro: Any,
    sueldo_nominal: float,
    salario_diario: float | None = None,
    dias_trabajados: float | None = None,
    dias_mes: float | None = None,
    fondo_ahorro_excel: float = 0.0,
    factor: float = DEFAULT_FONDO_AHORRO_FACTOR,
    uma_diaria: float = DEFAULT_UMA_DIARIA,
    tope_mode: str = DEFAULT_FONDO_AHORRO_TOPE_MODE,
) -> dict[str, Any]:
    del salario_diario  # Reserved for future client-validated variants.
    usa_fondo_ahorro = _normalize_yes_no(cuenta_fondo_ahorro)
    base_calc = sueldo_nominal * factor if usa_fondo_ahorro else 0.0
    cap_result = _calculate_fondo_ahorro_cap(
        uma_diaria=uma_diaria,
        tope_mode=tope_mode,
        dias_mes=dias_mes,
        dias_trabajados=dias_trabajados,
    )
    cap = cap_result["tope"]
    fondo_calc = base_calc if cap is None else min(base_calc, cap)
    diferencia = fondo_calc - fondo_ahorro_excel
    requiere_validacion = abs(diferencia) > 0.03 or bool(cap_result["advertencia"])

    if requiere_validacion:
        mensaje = (
            "Fondo de ahorro requiere validacion contra el criterio parametrizado. "
            f"Excel={round(fondo_ahorro_excel, 2)}, Calculado={round(fondo_calc, 2)}, "
            f"Diferencia={round(diferencia, 2)}, Factor={factor}, UMA={uma_diaria}, "
            f"TopeMode={cap_result['tope_mode_usado']}."
        )
        if cap_result["advertencia"]:
            mensaje += f" {cap_result['advertencia']}"
    else:
        mensaje = (
            "Fondo de ahorro coincide con el criterio parametrizado dentro de la tolerancia."
        )

    return {
        "fondo_ahorro_excel": round(fondo_ahorro_excel, 2),
        "fondo_ahorro_base_calc": round(base_calc, 2),
        "fondo_ahorro_tope_calc": cap,
        "fondo_ahorro_calc": round(fondo_calc, 2),
        "fondo_ahorro_diferencia": round(diferencia, 2),
        "fondo_ahorro_factor": factor,
        "uma_diaria": uma_diaria,
        "fondo_ahorro_tope_mode": cap_result["tope_mode_usado"],
        "usa_fondo_ahorro": usa_fondo_ahorro,
        "requiere_validacion": requiere_validacion,
        "mensaje_validacion": mensaje,
    }


def _calculate_sdi(salario_diario: float, factor_integracion_imss: float) -> float:
    return salario_diario * factor_integracion_imss


def _make_unique_columns(columns: list[Any]) -> list[str]:
    seen: dict[str, int] = {}
    result: list[str] = []
    for idx, column in enumerate(columns, start=1):
        name = "" if pd.isna(column) else str(column).strip()
        if not name:
            name = f"columna_{idx}"
        count = seen.get(name, 0)
        seen[name] = count + 1
        result.append(name if count == 0 else f"{name}_{count + 1}")
    return result


def _detect_header_row(excel_path: Path, sheet_name: str | int = 0) -> int:
    preview = pd.read_excel(excel_path, sheet_name=sheet_name, header=None, nrows=20)
    best_index = 0
    best_score = -1
    known = set(COLUMN_ALIASES) | set(ADMIN_BASE_ALIASES)
    for idx, row in preview.iterrows():
        normalized_cells = {_normalize_text(value) for value in row.tolist()}
        score = len(normalized_cells & known)
        if score > best_score:
            best_index = int(idx)
            best_score = score
    return best_index if best_score >= 2 else 0


def _read_excel(excel_path: Path) -> pd.DataFrame:
    header_row = _detect_header_row(excel_path)
    df = pd.read_excel(excel_path, sheet_name=0, header=header_row)
    df.columns = _make_unique_columns(df.columns.tolist())
    df = df.dropna(how="all").reset_index(drop=True)
    return df


def _read_excel_with_header(excel_path: Path) -> tuple[pd.DataFrame, int]:
    header_row = _detect_header_row(excel_path)
    df = pd.read_excel(excel_path, sheet_name=0, header=header_row)
    df.columns = _make_unique_columns(df.columns.tolist())
    df = df.dropna(how="all").reset_index(drop=True)
    return df, header_row


def _canonicalize_columns(df: pd.DataFrame) -> tuple[pd.DataFrame, dict[str, str]]:
    canonical = df.copy()
    mapping: dict[str, str] = {}
    sources: dict[str, list[str]] = {}

    for original in canonical.columns:
        normalized = _normalize_text(original)
        target = COLUMN_ALIASES.get(normalized)
        if target:
            sources.setdefault(target, []).append(original)

    for target, source_columns in sources.items():
        preferred = sorted(
            source_columns,
            key=lambda source: 0 if _normalize_text(source) == _normalize_text(target) else 1,
        )
        combined = canonical[preferred[0]].copy()
        for source in preferred[1:]:
            blank_mask = combined.apply(_is_blank)
            combined.loc[blank_mask] = canonical.loc[blank_mask, source]
        canonical[target] = combined
        mapping[target] = preferred[0]

    return canonical, mapping


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    return str(value).strip() == ""


def _add_issue(
    issues: list[dict[str, Any]],
    severidad: str,
    tipo: str,
    mensaje: str,
    fila: int | None = None,
    columna: str | None = None,
) -> None:
    issues.append(
        {
            "severidad": severidad,
            "tipo": tipo,
            "fila": fila,
            "columna": columna,
            "mensaje": mensaje,
        }
    )


def _to_numeric(
    series: pd.Series,
    column: str,
    issues: list[dict[str, Any]],
    *,
    required: bool,
) -> pd.Series:
    converted = pd.to_numeric(series, errors="coerce")
    invalid_mask = converted.isna()

    for index in series[invalid_mask].index:
        value = series.loc[index]
        if _is_blank(value) and not required:
            _add_issue(
                issues,
                "advertencia",
                "valor_vacio",
                f"'{column}' esta vacio; se tomara como 0.",
                fila=int(index) + 2,
                columna=column,
            )
            converted.loc[index] = 0
            continue

        _add_issue(
            issues,
            "critico",
            "numero_invalido",
            f"'{column}' debe ser numerico.",
            fila=int(index) + 2,
            columna=column,
        )

    return converted


def _build_summary_sheet(
    ok: bool,
    empleados_procesados: int,
    total_prenomina: float | None,
    critical_count: int,
    warning_count: int,
) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"metrica": "ok", "valor": ok},
            {"metrica": "empleados_procesados", "valor": empleados_procesados},
            {"metrica": "total_prenomina", "valor": total_prenomina},
            {"metrica": "errores_criticos", "valor": critical_count},
            {"metrica": "advertencias", "valor": warning_count},
        ]
    )


def _format_workbook(path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)
            for cell in column_cells:
                value = cell.value
                if value is not None:
                    max_length = max(max_length, len(str(value)))
            worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 42)

    workbook.save(path)


def _clean_unique_suffix(column: Any) -> str:
    return re.sub(r"\s+\d+$", "", _normalize_text(column))


def _build_admin_column_map(df: pd.DataFrame) -> dict[str, dict[str, Any]]:
    mapping: dict[str, dict[str, Any]] = {}
    occurrences: dict[str, int] = {}

    for idx, column in enumerate(df.columns):
        normalized = _clean_unique_suffix(column)
        target = ADMIN_BASE_ALIASES.get(normalized)
        if not target:
            continue

        if target in ["sueldo_bruto_mensual", "sueldo_bruto_quincenal"]:
            occurrences[target] = occurrences.get(target, 0) + 1
            if target == "sueldo_bruto_mensual" and occurrences[target] == 2:
                target = "sueldo_bruto_mensual_final"
            elif target == "sueldo_bruto_quincenal" and occurrences[target] == 2:
                target = "sueldo_bruto_quincenal_final"

        if target not in mapping:
            mapping[target] = {"name": column, "index": idx}

    return mapping


def _is_admin_prenomina(df: pd.DataFrame) -> bool:
    mapping = _build_admin_column_map(df)
    signals = {
        "sueldo_nominal",
        "percepcion_sueldos",
        "honorarios_asimilados",
        "sueldo_bruto_mensual",
        "sueldo_bruto_quincenal",
        "observaciones",
    }
    return len(signals & set(mapping)) >= 4


def _safe_number(value: Any) -> float:
    if _is_blank(value):
        return 0.0
    converted = pd.to_numeric(pd.Series([value]), errors="coerce").iloc[0]
    if pd.isna(converted):
        return 0.0
    return float(converted)


def _row_value(row: pd.Series, mapping: dict[str, dict[str, Any]], key: str) -> Any:
    if key not in mapping:
        return None
    return row[mapping[key]["name"]]


def _row_number(row: pd.Series, mapping: dict[str, dict[str, Any]], key: str) -> float:
    return _safe_number(_row_value(row, mapping, key))


def _parse_admin_observaciones(observaciones: Any) -> dict[str, Any]:
    text = "" if _is_blank(observaciones) else str(observaciones).strip()
    normalized = _normalize_text(text)
    tags: list[str] = []
    warnings: list[str] = []
    descuento_dias = 0

    day_match = re.search(r"descontar\s+(\d+(?:\.\d+)?)\s+dia", normalized)
    if day_match:
        descuento_dias = int(float(day_match.group(1)))
        tags.append("descuento_dias")

    if normalized in {"", "ok"} or normalized.startswith("ok "):
        tags.append("ok")

    if "prestamo" in normalized or "adeudo" in normalized:
        tags.append("prestamo_adeudo")
        warnings.append("Observacion de prestamo/adeudo requiere confirmacion humana.")

    if "factura" in normalized or "facturado" in normalized:
        tags.append("factura")

    if "17" in normalized:
        tags.append("pago_dia_17")

    if "15" in normalized:
        tags.append("pago_dia_15")

    if "comision" in normalized:
        tags.append("comisiones")

    if "ultimo dia" in normalized or "ultimo dia" in normalized:
        tags.append("pago_ultimo_dia")

    if "fondo de ahorro" in normalized:
        tags.append("fondo_ahorro")
        warnings.append("Observacion relacionada con fondo de ahorro requiere revision.")

    if text and tags == []:
        tags.append("observacion_no_clasificada")
        warnings.append("Observacion no clasificada automaticamente.")

    return {
        "texto": text,
        "tags": ", ".join(dict.fromkeys(tags)),
        "descuento_dias": descuento_dias,
        "warnings": warnings,
    }


def _external_links_summary(excel_path: Path) -> list[str]:
    from openpyxl import load_workbook

    workbook = load_workbook(excel_path, read_only=True, keep_links=True)
    links: list[str] = []
    for link in getattr(workbook, "_external_links", []):
        file_link = getattr(link, "file_link", None)
        target = getattr(file_link, "Target", None)
        if target:
            links.append(str(target))
    workbook.close()
    return links


def _formula_for_cell(excel_path: Path, sheet_name: str, row: int, col_index: int) -> str | None:
    from openpyxl import load_workbook

    workbook = load_workbook(excel_path, data_only=False)
    worksheet = workbook[sheet_name]
    cell = worksheet.cell(row=row, column=col_index + 1)
    formula = cell.value if cell.data_type == "f" else None
    workbook.close()
    return formula


def _admin_sheet_value(worksheet: Any, row: int, mapping: dict[str, dict[str, Any]], key: str) -> Any:
    if key not in mapping:
        return None
    return worksheet.cell(row=row, column=int(mapping[key]["index"]) + 1).value


def _format_identifier(value: Any, fallback: str) -> str:
    if _is_blank(value):
        return fallback
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _has_external_reference(formula: Any) -> bool:
    if _is_blank(formula):
        return False
    normalized = _normalize_text(formula)
    text = str(formula)
    return any(
        marker in text or marker in normalized
        for marker in ["[", "]", ".xlsx", "UMA Y SMG", "uma y smg"]
    )


def _load_admin_prenomina_config(issues: list[dict[str, Any]]) -> dict[str, Any]:
    use_excel_fondo_ahorro = _env_bool(
        "USE_EXCEL_FONDO_AHORRO", DEFAULT_USE_EXCEL_FONDO_AHORRO
    )
    config: dict[str, Any] = {
        "use_excel_fondo_ahorro": use_excel_fondo_ahorro,
        "fondo_ahorro_factor": DEFAULT_FONDO_AHORRO_FACTOR,
        "uma_diaria": DEFAULT_UMA_DIARIA,
        "fondo_ahorro_tope_mode": _normalize_text(
            os.getenv("FONDO_AHORRO_TOPE_MODE", DEFAULT_FONDO_AHORRO_TOPE_MODE)
        )
        or DEFAULT_FONDO_AHORRO_TOPE_MODE,
        "dias_base_periodo": DEFAULT_DIAS_BASE_PERIODO,
        "dias_mes": DEFAULT_DIAS_MES,
    }

    for env_name, config_key, default in [
        ("FONDO_AHORRO_FACTOR", "fondo_ahorro_factor", DEFAULT_FONDO_AHORRO_FACTOR),
        ("UMA_DIARIA", "uma_diaria", DEFAULT_UMA_DIARIA),
        ("DIAS_BASE_PERIODO", "dias_base_periodo", DEFAULT_DIAS_BASE_PERIODO),
        ("DIAS_MES", "dias_mes", DEFAULT_DIAS_MES),
    ]:
        try:
            config[config_key] = _env_float(env_name, default)
        except ValueError:
            severity = "critico" if not use_excel_fondo_ahorro else "advertencia"
            _add_issue(
                issues,
                severity,
                "variable_entorno_invalida",
                f"{env_name} no es numerica; se usa default {default}.",
                columna=env_name,
            )

    if config["fondo_ahorro_factor"] <= 0:
        severity = "critico" if not use_excel_fondo_ahorro else "advertencia"
        _add_issue(
            issues,
            severity,
            "fondo_ahorro_factor_invalido",
            "FONDO_AHORRO_FACTOR debe ser mayor a 0.",
            columna="FONDO_AHORRO_FACTOR",
        )
        config["fondo_ahorro_factor"] = DEFAULT_FONDO_AHORRO_FACTOR

    if config["uma_diaria"] <= 0:
        severity = "critico" if not use_excel_fondo_ahorro else "advertencia"
        _add_issue(
            issues,
            severity,
            "uma_diaria_invalida",
            "UMA_DIARIA debe ser mayor a 0.",
            columna="UMA_DIARIA",
        )
        config["uma_diaria"] = DEFAULT_UMA_DIARIA

    if config["fondo_ahorro_tope_mode"] not in FONDO_AHORRO_TOPE_MODES:
        _add_issue(
            issues,
            "advertencia",
            "fondo_ahorro_tope_mode_invalido",
            (
                "FONDO_AHORRO_TOPE_MODE debe ser none, mensual, quincenal o "
                f"proporcional; se usa {DEFAULT_FONDO_AHORRO_TOPE_MODE}."
            ),
            columna="FONDO_AHORRO_TOPE_MODE",
        )
        config["fondo_ahorro_tope_mode"] = DEFAULT_FONDO_AHORRO_TOPE_MODE

    return config


def procesar_prenomina_administrativa(
    excel_path: Path,
    result_path: Path,
    base_df: pd.DataFrame,
    header_row: int,
) -> dict[str, Any]:
    from openpyxl import load_workbook

    issues: list[dict[str, Any]] = []
    differences: list[dict[str, Any]] = []
    mapping = _build_admin_column_map(base_df)
    config = _load_admin_prenomina_config(issues)

    for column in ADMIN_REQUIRED_COLUMNS:
        if column not in mapping:
            _add_issue(
                issues,
                "critico",
                "columna_faltante",
                f"Falta la columna requerida '{column}' para prenomina administrativa.",
                columna=column,
            )

    for link in _external_links_summary(excel_path):
        _add_issue(
            issues,
            "advertencia",
            "dependencia_externa",
            "El archivo contiene una referencia externa. Conviene reemplazarla por una variable local.",
            columna=link,
        )

    if "nombre_completo" in mapping:
        employee_mask = base_df[mapping["nombre_completo"]["name"]].apply(lambda value: not _is_blank(value))
        employee_df = base_df[employee_mask].copy()
    else:
        employee_df = pd.DataFrame()

    critical_count = sum(issue["severidad"] == "critico" for issue in issues)
    if critical_count:
        prenomina_df = pd.DataFrame(
            [
                {
                    "estado": "No calculada",
                    "motivo": "Hay columnas criticas faltantes para prenomina administrativa.",
                }
            ]
        )
        inconsistencias_df = pd.DataFrame(issues)
        resumen_df = _build_summary_sheet(False, 0, None, critical_count, 0)
        with pd.ExcelWriter(result_path, engine="openpyxl") as writer:
            base_df.to_excel(writer, index=False, sheet_name="Base original")
            prenomina_df.to_excel(writer, index=False, sheet_name="Prenomina")
            inconsistencias_df.to_excel(writer, index=False, sheet_name="Inconsistencias")
            resumen_df.to_excel(writer, index=False, sheet_name="Resumen")
        _format_workbook(result_path)
        return {
            "ok": False,
            "mensaje": "Se detectaron columnas criticas faltantes.",
            "archivo_generado": str(result_path),
            "empleados_procesados": 0,
            "total_prenomina": None,
            "inconsistencias": issues,
            "columnas_detectadas": {key: value["name"] for key, value in mapping.items()},
            "tipo_archivo": "prenomina_administrativa",
        }

    workbook_formula = load_workbook(excel_path, data_only=False)
    worksheet_formula = workbook_formula[workbook_formula.sheetnames[0]]

    rows: list[dict[str, Any]] = []
    tolerance = 0.03
    comparison_fields = [
        ("percepcion_sueldos", "percepcion_sueldos_calc", "Percepcion sueldos"),
        ("sueldo_bruto_mensual", "sueldo_bruto_mensual_calc", "Sueldo bruto mensual"),
        ("sueldo_bruto_quincenal", "sueldo_bruto_quincenal_calc", "Sueldo bruto quincenal"),
        ("sueldo_bruto_mensual_final", "sueldo_bruto_mensual_final_calc", "Sueldo bruto mensual final"),
        ("sueldo_bruto_quincenal_final", "sueldo_bruto_quincenal_final_calc", "Sueldo bruto quincenal final"),
    ]

    for index, row in employee_df.iterrows():
        excel_row = int(index) + header_row + 2
        numero = _admin_sheet_value(worksheet_formula, excel_row, mapping, "numero")
        codigo = _admin_sheet_value(worksheet_formula, excel_row, mapping, "codigo")
        nombre = str(_admin_sheet_value(worksheet_formula, excel_row, mapping, "nombre_completo")).strip()
        empleado_id = _format_identifier(codigo, f"fila_{excel_row}")

        if _is_blank(codigo):
            _add_issue(
                issues,
                "advertencia",
                "codigo_vacio",
                "El codigo del empleado esta vacio; se uso la fila como identificador temporal.",
                fila=excel_row,
                columna="Cod.",
            )

        observaciones = _admin_sheet_value(worksheet_formula, excel_row, mapping, "observaciones")
        parsed_obs = _parse_admin_observaciones(observaciones)
        for warning in parsed_obs["warnings"]:
            _add_issue(
                issues,
                "advertencia",
                "observacion_requiere_revision",
                warning,
                fila=excel_row,
                columna="OBSERVACIONES",
            )

        sueldo_nominal = _row_number(row, mapping, "sueldo_nominal")
        puntualidad = _row_number(row, mapping, "puntualidad")
        asistencia = _row_number(row, mapping, "asistencia")
        vales_despensa = _row_number(row, mapping, "vales_despensa")
        fondo_ahorro_excel = _row_number(row, mapping, "fondo_ahorro")
        honorarios_asimilados = _row_number(row, mapping, "honorarios_asimilados")
        gasolina_sueldo = _row_number(row, mapping, "gasolina_sueldo")
        socio = _row_number(row, mapping, "socio")
        efectivo = _row_number(row, mapping, "efectivo")
        facturado = _row_number(row, mapping, "facturado")
        deuda_carro = _row_number(row, mapping, "deuda_carro")
        descuento = _row_number(row, mapping, "descuento")
        salario_diario = _row_number(row, mapping, "salario_diario")
        factor_integracion_imss = _row_number(row, mapping, "factor_integracion_imss")
        salario_diario_integrado_excel = _row_number(row, mapping, "salario_diario_integrado")
        salario_diario_integrado_calc = _calculate_sdi(
            salario_diario, factor_integracion_imss
        )
        salario_diario_integrado_diferencia = (
            salario_diario_integrado_calc - salario_diario_integrado_excel
        )
        dias_trabajados = (
            _row_number(row, mapping, "dias_trabajados")
            if "dias_trabajados" in mapping
            else None
        )
        dias_mes = _row_number(row, mapping, "dias_mes") if "dias_mes" in mapping else config["dias_mes"]

        fondo_formula = None
        if "fondo_ahorro" in mapping:
            fondo_cell = worksheet_formula.cell(
                row=excel_row,
                column=int(mapping["fondo_ahorro"]["index"]) + 1,
            )
            fondo_formula = fondo_cell.value if fondo_cell.data_type == "f" else None
            if _has_external_reference(fondo_formula):
                _add_issue(
                    issues,
                    "advertencia",
                    "formula_externa_fondo_ahorro",
                    (
                        "La celda de Fondo de ahorro contiene una formula con referencia externa. "
                        "El sistema no puede leer el archivo externo; sustituye el factor con "
                        "FONDO_AHORRO_FACTOR y reporta diferencia para validacion."
                    ),
                    fila=excel_row,
                    columna=mapping["fondo_ahorro"]["name"],
                )

        fondo_ahorro_result = _calculate_fondo_ahorro(
            cuenta_fondo_ahorro=_admin_sheet_value(
                worksheet_formula, excel_row, mapping, "cuenta_fondo_ahorro"
            ),
            sueldo_nominal=sueldo_nominal,
            salario_diario=salario_diario,
            dias_trabajados=dias_trabajados,
            dias_mes=dias_mes,
            fondo_ahorro_excel=fondo_ahorro_excel,
            factor=config["fondo_ahorro_factor"],
            uma_diaria=config["uma_diaria"],
            tope_mode=config["fondo_ahorro_tope_mode"],
        )
        if abs(fondo_ahorro_result["fondo_ahorro_diferencia"]) > tolerance:
            differences.append(
                {
                    "fila_excel": excel_row,
                    "empleado_id": empleado_id,
                    "nombre": nombre,
                    "concepto": "Fondo de ahorro",
                    "valor_excel": fondo_ahorro_result["fondo_ahorro_excel"],
                    "valor_calculado": fondo_ahorro_result["fondo_ahorro_calc"],
                    "diferencia": fondo_ahorro_result["fondo_ahorro_diferencia"],
                    "formula_excel": fondo_formula,
                }
            )
            _add_issue(
                issues,
                "advertencia",
                "diferencia_fondo_ahorro",
                (
                    "Fondo de ahorro difiere del calculo parametrizado. "
                    f"Excel={fondo_ahorro_result['fondo_ahorro_excel']}, "
                    f"Calculado={fondo_ahorro_result['fondo_ahorro_calc']}, "
                    f"Diferencia={fondo_ahorro_result['fondo_ahorro_diferencia']}, "
                    f"Factor={fondo_ahorro_result['fondo_ahorro_factor']}, "
                    f"UMA={fondo_ahorro_result['uma_diaria']}, "
                    f"TopeMode={fondo_ahorro_result['fondo_ahorro_tope_mode']}. "
                    "Validar con usuario."
                ),
                fila=excel_row,
                columna=mapping["fondo_ahorro"]["name"],
            )

        if abs(salario_diario_integrado_diferencia) > tolerance:
            _add_issue(
                issues,
                "advertencia",
                "diferencia_sdi",
                (
                    "Salario diario integrado difiere del calculo salario_diario x "
                    f"factor_integracion_imss. Excel={round(salario_diario_integrado_excel, 2)}, "
                    f"Calculado={round(salario_diario_integrado_calc, 2)}, "
                    f"Diferencia={round(salario_diario_integrado_diferencia, 2)}."
                ),
                fila=excel_row,
                columna=mapping.get("salario_diario_integrado", {}).get("name"),
            )

        puntualidad_validacion_10_sdi = (
            salario_diario_integrado_calc * 0.10 * config["dias_base_periodo"]
        )
        asistencia_validacion_10_sdi = (
            salario_diario_integrado_calc * 0.10 * config["dias_base_periodo"]
        )
        puntualidad_validacion_diferencia = puntualidad_validacion_10_sdi - puntualidad
        asistencia_validacion_diferencia = asistencia_validacion_10_sdi - asistencia
        if (
            abs(puntualidad_validacion_diferencia) > tolerance
            or abs(asistencia_validacion_diferencia) > tolerance
        ):
            _add_issue(
                issues,
                "advertencia",
                "validacion_puntualidad_asistencia",
                (
                    "Validacion exploratoria: puntualidad/asistencia contra 10% de SDI "
                    f"por DIAS_BASE_PERIODO={config['dias_base_periodo']}. "
                    f"Puntualidad Excel={round(puntualidad, 2)}, "
                    f"Calculada={round(puntualidad_validacion_10_sdi, 2)}, "
                    f"Diferencia={round(puntualidad_validacion_diferencia, 2)}. "
                    f"Asistencia Excel={round(asistencia, 2)}, "
                    f"Calculada={round(asistencia_validacion_10_sdi, 2)}, "
                    f"Diferencia={round(asistencia_validacion_diferencia, 2)}."
                ),
                fila=excel_row,
                columna="PUNTUALIDAD / ASISTENCIA",
            )

        fondo_ahorro_para_calculo = (
            fondo_ahorro_excel
            if config["use_excel_fondo_ahorro"]
            else fondo_ahorro_result["fondo_ahorro_calc"]
        )
        fondo_ahorro_fuente_usada = (
            "excel" if config["use_excel_fondo_ahorro"] else "calculado"
        )

        percepcion_sueldos_calc = (
            sueldo_nominal
            + puntualidad
            + asistencia
            + vales_despensa
            + fondo_ahorro_para_calculo
        )
        sueldo_bruto_mensual_calc = (
            percepcion_sueldos_calc
            + honorarios_asimilados
            + gasolina_sueldo
            + socio
            + efectivo
            + facturado
            + deuda_carro
        )
        sueldo_bruto_quincenal_calc = sueldo_bruto_mensual_calc / 2
        sueldo_bruto_mensual_final_calc = sueldo_bruto_mensual_calc - descuento
        sueldo_bruto_quincenal_final_calc = sueldo_bruto_mensual_final_calc / 2

        descuento_dias = parsed_obs["descuento_dias"]
        descuento_por_dias = 0.0
        if descuento_dias:
            quincena_base = sueldo_bruto_mensual_final_calc / 2
            descuento_por_dias = quincena_base / 15 * descuento_dias
            sueldo_bruto_quincenal_final_calc = quincena_base - descuento_por_dias

        output_row = {
            "fila_excel": excel_row,
            "numero": numero,
            "empleado_id": empleado_id,
            "codigo": codigo,
            "empresa": _admin_sheet_value(worksheet_formula, excel_row, mapping, "empresa"),
            "nombre": nombre,
            "cuenta_fondo_ahorro": _admin_sheet_value(
                worksheet_formula, excel_row, mapping, "cuenta_fondo_ahorro"
            ),
            "salario_diario": round(salario_diario, 2),
            "factor_integracion_imss": round(factor_integracion_imss, 6),
            "salario_diario_integrado_excel": round(salario_diario_integrado_excel, 2),
            "salario_diario_integrado_calc": round(salario_diario_integrado_calc, 2),
            "salario_diario_integrado_diferencia": round(
                salario_diario_integrado_diferencia, 2
            ),
            "sueldo_nominal": round(sueldo_nominal, 2),
            "puntualidad": round(puntualidad, 2),
            "puntualidad_validacion_10_sdi": round(puntualidad_validacion_10_sdi, 2),
            "puntualidad_validacion_diferencia": round(
                puntualidad_validacion_diferencia, 2
            ),
            "asistencia": round(asistencia, 2),
            "asistencia_validacion_10_sdi": round(asistencia_validacion_10_sdi, 2),
            "asistencia_validacion_diferencia": round(asistencia_validacion_diferencia, 2),
            "vales_despensa": round(vales_despensa, 2),
            "fondo_ahorro": round(fondo_ahorro_para_calculo, 2),
            "fondo_ahorro_excel": fondo_ahorro_result["fondo_ahorro_excel"],
            "fondo_ahorro_base_calc": fondo_ahorro_result["fondo_ahorro_base_calc"],
            "fondo_ahorro_tope_calc": fondo_ahorro_result["fondo_ahorro_tope_calc"],
            "fondo_ahorro_calc": fondo_ahorro_result["fondo_ahorro_calc"],
            "fondo_ahorro_diferencia": fondo_ahorro_result["fondo_ahorro_diferencia"],
            "fondo_ahorro_factor": fondo_ahorro_result["fondo_ahorro_factor"],
            "uma_diaria": fondo_ahorro_result["uma_diaria"],
            "fondo_ahorro_tope_mode": fondo_ahorro_result["fondo_ahorro_tope_mode"],
            "fondo_ahorro_fuente_usada": fondo_ahorro_fuente_usada,
            "fondo_ahorro_requiere_validacion": fondo_ahorro_result[
                "requiere_validacion"
            ],
            "percepcion_sueldos_calc": round(percepcion_sueldos_calc, 2),
            "honorarios_asimilados": round(honorarios_asimilados, 2),
            "gasolina_sueldo": round(gasolina_sueldo, 2),
            "socio": round(socio, 2),
            "efectivo": round(efectivo, 2),
            "facturado": round(facturado, 2),
            "deuda_carro": round(deuda_carro, 2),
            "sueldo_bruto_mensual_calc": round(sueldo_bruto_mensual_calc, 2),
            "sueldo_bruto_quincenal_calc": round(sueldo_bruto_quincenal_calc, 2),
            "descuento": round(descuento, 2),
            "descuento_dias": descuento_dias,
            "descuento_por_dias": round(descuento_por_dias, 2),
            "sueldo_bruto_mensual_final_calc": round(sueldo_bruto_mensual_final_calc, 2),
            "sueldo_bruto_quincenal_final_calc": round(sueldo_bruto_quincenal_final_calc, 2),
            "observaciones": parsed_obs["texto"],
            "incidencias_detectadas": parsed_obs["tags"],
        }

        for excel_key, calc_key, label in comparison_fields:
            excel_value = _row_number(row, mapping, excel_key)
            calc_value = float(output_row[calc_key])
            difference = round(calc_value - excel_value, 2)
            if abs(difference) > tolerance:
                col_index = int(mapping[excel_key]["index"])
                formula_cell = worksheet_formula.cell(row=excel_row, column=col_index + 1)
                formula = formula_cell.value if formula_cell.data_type == "f" else None
                differences.append(
                    {
                        "fila_excel": excel_row,
                        "empleado_id": empleado_id,
                        "nombre": nombre,
                        "concepto": label,
                        "valor_excel": round(excel_value, 2),
                        "valor_calculado": round(calc_value, 2),
                        "diferencia": difference,
                        "formula_excel": formula,
                    }
                )
                _add_issue(
                    issues,
                    "advertencia",
                    "diferencia_formula",
                    f"{label}: diferencia entre Excel y calculo Python de {difference}.",
                    fila=excel_row,
                    columna=mapping[excel_key]["name"],
                )

        rows.append(output_row)

    workbook_formula.close()

    prenomina_df = pd.DataFrame(rows)
    differences_df = pd.DataFrame(differences)
    if differences_df.empty:
        differences_df = pd.DataFrame(
            columns=[
                "fila_excel",
                "empleado_id",
                "nombre",
                "concepto",
                "valor_excel",
                "valor_calculado",
                "diferencia",
                "formula_excel",
            ]
        )

    inconsistencias_df = pd.DataFrame(issues)
    if inconsistencias_df.empty:
        inconsistencias_df = pd.DataFrame(
            columns=["severidad", "tipo", "fila", "columna", "mensaje"]
        )

    critical_count = sum(issue["severidad"] == "critico" for issue in issues)
    warning_count = sum(issue["severidad"] == "advertencia" for issue in issues)
    empleados_procesados = int(len(prenomina_df))
    total_prenomina = (
        float(round(prenomina_df["sueldo_bruto_quincenal_final_calc"].sum(), 2))
        if empleados_procesados and "sueldo_bruto_quincenal_final_calc" in prenomina_df
        else 0.0
    )

    def _sum_prenomina_column(column: str) -> float:
        if column not in prenomina_df:
            return 0.0
        return float(round(prenomina_df[column].sum(), 2))

    resumen_rows = [
        {"metrica": "tipo_archivo", "valor": "prenomina_administrativa"},
        {"metrica": "ok", "valor": critical_count == 0},
        {"metrica": "empleados_procesados", "valor": empleados_procesados},
        {"metrica": "total_quincenal_final", "valor": total_prenomina},
        {"metrica": "total_percepcion_sueldos", "valor": _sum_prenomina_column("percepcion_sueldos_calc")},
        {"metrica": "total_honorarios_asimilados", "valor": _sum_prenomina_column("honorarios_asimilados")},
        {"metrica": "total_gasolina", "valor": _sum_prenomina_column("gasolina_sueldo")},
        {"metrica": "total_socio", "valor": _sum_prenomina_column("socio")},
        {"metrica": "total_efectivo", "valor": _sum_prenomina_column("efectivo")},
        {"metrica": "total_facturado", "valor": _sum_prenomina_column("facturado")},
        {"metrica": "total_deuda_carro", "valor": _sum_prenomina_column("deuda_carro")},
        {"metrica": "fondo_ahorro_factor", "valor": config["fondo_ahorro_factor"]},
        {"metrica": "uma_diaria", "valor": config["uma_diaria"]},
        {"metrica": "fondo_ahorro_tope_mode", "valor": config["fondo_ahorro_tope_mode"]},
        {"metrica": "use_excel_fondo_ahorro", "valor": config["use_excel_fondo_ahorro"]},
        {"metrica": "total_fondo_ahorro_excel", "valor": _sum_prenomina_column("fondo_ahorro_excel")},
        {"metrica": "total_fondo_ahorro_calc", "valor": _sum_prenomina_column("fondo_ahorro_calc")},
        {"metrica": "total_diferencia_fondo_ahorro", "valor": _sum_prenomina_column("fondo_ahorro_diferencia")},
        {
            "metrica": "empleados_con_fondo_ahorro",
            "valor": int(prenomina_df["fondo_ahorro_calc"].gt(0).sum())
            if "fondo_ahorro_calc" in prenomina_df
            else 0,
        },
        {
            "metrica": "empleados_fondo_ahorro_con_diferencia",
            "valor": int(prenomina_df["fondo_ahorro_diferencia"].abs().gt(tolerance).sum())
            if "fondo_ahorro_diferencia" in prenomina_df
            else 0,
        },
        {
            "metrica": "requiere_validacion_fondo_ahorro",
            "valor": bool(prenomina_df["fondo_ahorro_requiere_validacion"].any())
            if "fondo_ahorro_requiere_validacion" in prenomina_df
            else False,
        },
        {"metrica": "errores_criticos", "valor": critical_count},
        {"metrica": "advertencias", "valor": warning_count},
        {"metrica": "diferencias_con_excel", "valor": len(differences_df)},
    ]
    resumen_df = pd.DataFrame(resumen_rows)

    with pd.ExcelWriter(result_path, engine="openpyxl") as writer:
        base_df.to_excel(writer, index=False, sheet_name="Base original")
        prenomina_df.to_excel(writer, index=False, sheet_name="Prenomina")
        inconsistencias_df.to_excel(writer, index=False, sheet_name="Inconsistencias")
        differences_df.to_excel(writer, index=False, sheet_name="Diferencias")
        resumen_df.to_excel(writer, index=False, sheet_name="Resumen")

    _format_workbook(result_path)

    mensaje = "Prenomina administrativa calculada correctamente."
    if warning_count:
        mensaje += " Revisa las advertencias antes de usarla como calculo final."

    return {
        "ok": critical_count == 0,
        "mensaje": mensaje,
        "archivo_generado": str(result_path),
        "empleados_procesados": empleados_procesados,
        "total_prenomina": total_prenomina,
        "inconsistencias": issues,
        "columnas_detectadas": {key: value["name"] for key, value in mapping.items()},
        "tipo_archivo": "prenomina_administrativa",
    }


def procesar_prenomina_excel(input_path: str, output_path: str | None = None) -> dict[str, Any]:
    """Calcula la prenomina operativa desde Excel y genera el reporte de salida."""
    excel_path = Path(input_path)
    result_path = Path(output_path) if output_path else DEFAULT_OUTPUT_PATH
    issues: list[dict[str, Any]] = []

    if not excel_path.exists():
        raise FileNotFoundError(f"No existe el archivo: {excel_path}")

    result_path.parent.mkdir(parents=True, exist_ok=True)
    base_df, header_row = _read_excel_with_header(excel_path)
    work_df, mapping = _canonicalize_columns(base_df)
    has_attendance_columns = all(column in work_df.columns for column in REQUIRED_CALC_COLUMNS)
    if not has_attendance_columns and _is_admin_prenomina(base_df):
        return procesar_prenomina_administrativa(
            excel_path=excel_path,
            result_path=result_path,
            base_df=base_df,
            header_row=header_row,
        )

    for column in REQUIRED_CALC_COLUMNS:
        if column not in work_df.columns:
            _add_issue(
                issues,
                "critico",
                "columna_faltante",
                f"Falta la columna requerida '{column}'.",
                columna=column,
            )

    if not any(issue["severidad"] == "critico" for issue in issues):
        work_df["empleado_id"] = work_df["empleado_id"].astype(str).str.strip()
        work_df["nombre"] = work_df["nombre"].astype(str).str.strip()

        for index, value in work_df["empleado_id"].items():
            if _is_blank(value) or value.lower() == "nan":
                _add_issue(
                    issues,
                    "critico",
                    "empleado_id_vacio",
                    "El empleado_id no puede estar vacio.",
                    fila=int(index) + 2,
                    columna="empleado_id",
                )

        for index, value in work_df["nombre"].items():
            if _is_blank(value) or value.lower() == "nan":
                _add_issue(
                    issues,
                    "critico",
                    "nombre_vacio",
                    "El nombre no puede estar vacio.",
                    fila=int(index) + 2,
                    columna="nombre",
                )

        work_df["fecha"] = pd.to_datetime(work_df["fecha"], errors="coerce")
        for index in work_df[work_df["fecha"].isna()].index:
            _add_issue(
                issues,
                "critico",
                "fecha_invalida",
                "La fecha no es valida.",
                fila=int(index) + 2,
                columna="fecha",
            )

        for column in ["horas_trabajadas", "sueldo_hora", "bono", "descuento"]:
            work_df[column] = _to_numeric(
                work_df[column],
                column,
                issues,
                required=column in ["horas_trabajadas", "sueldo_hora"],
            )

        for column in ["horas_trabajadas", "sueldo_hora", "bono", "descuento"]:
            for index in work_df[work_df[column] < 0].index:
                _add_issue(
                    issues,
                    "critico",
                    "valor_negativo",
                    f"'{column}' no puede ser negativo.",
                    fila=int(index) + 2,
                    columna=column,
                )

        duplicate_mask = work_df.duplicated(subset=["empleado_id", "fecha"], keep=False)
        for index in work_df[duplicate_mask].index:
            _add_issue(
                issues,
                "critico",
                "registro_duplicado",
                "Registro duplicado por empleado_id + fecha.",
                fila=int(index) + 2,
                columna="empleado_id, fecha",
            )

    critical_count = sum(issue["severidad"] == "critico" for issue in issues)
    warning_count = sum(issue["severidad"] == "advertencia" for issue in issues)
    ok = critical_count == 0

    if ok:
        prenomina_df = (
            work_df.groupby(["empleado_id", "nombre"], dropna=False)
            .agg(
                total_horas=("horas_trabajadas", "sum"),
                sueldo_hora=("sueldo_hora", "first"),
                total_bonos=("bono", "sum"),
                total_descuentos=("descuento", "sum"),
            )
            .reset_index()
        )
        prenomina_df["subtotal"] = prenomina_df["total_horas"] * prenomina_df["sueldo_hora"]
        prenomina_df["total_pagar"] = (
            prenomina_df["subtotal"] + prenomina_df["total_bonos"] - prenomina_df["total_descuentos"]
        )
        amount_columns = ["sueldo_hora", "total_bonos", "total_descuentos", "subtotal", "total_pagar"]
        prenomina_df[amount_columns] = prenomina_df[amount_columns].round(2)
        empleados_procesados = int(prenomina_df["empleado_id"].nunique())
        total_prenomina = float(round(prenomina_df["total_pagar"].sum(), 2))
        mensaje = "Prenomina calculada correctamente."
    else:
        prenomina_df = pd.DataFrame(
            [
                {
                    "estado": "No calculada",
                    "motivo": "Hay errores criticos. Corrige las inconsistencias antes de generar el total final.",
                }
            ]
        )
        empleados_procesados = 0
        total_prenomina = None
        mensaje = "Se detectaron errores criticos. Se genero reporte de inconsistencias sin total final."

    inconsistencias_df = pd.DataFrame(issues)
    if inconsistencias_df.empty:
        inconsistencias_df = pd.DataFrame(
            columns=["severidad", "tipo", "fila", "columna", "mensaje"]
        )

    resumen_df = _build_summary_sheet(
        ok=ok,
        empleados_procesados=empleados_procesados,
        total_prenomina=total_prenomina,
        critical_count=critical_count,
        warning_count=warning_count,
    )

    with pd.ExcelWriter(result_path, engine="openpyxl") as writer:
        base_df.to_excel(writer, index=False, sheet_name="Base original")
        prenomina_df.to_excel(writer, index=False, sheet_name="Prenomina")
        inconsistencias_df.to_excel(writer, index=False, sheet_name="Inconsistencias")
        resumen_df.to_excel(writer, index=False, sheet_name="Resumen")

    _format_workbook(result_path)

    return {
        "ok": ok,
        "mensaje": mensaje,
        "archivo_generado": str(result_path),
        "empleados_procesados": empleados_procesados,
        "total_prenomina": total_prenomina,
        "inconsistencias": issues,
        "columnas_detectadas": mapping,
    }


@function_tool
def calcular_prenomina_excel(input_path: str, output_path: str = "") -> dict[str, Any]:
    """Lee un Excel de asistencia/incidencias, valida datos y genera la prenomina.

    Args:
        input_path: Ruta local del archivo Excel de entrada.
        output_path: Ruta local del archivo Excel de salida.
    """
    return procesar_prenomina_excel(input_path=input_path, output_path=output_path or None)
